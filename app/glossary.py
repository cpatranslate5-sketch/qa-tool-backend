"""
Parsing for the project's master glossary document: one sheet, first
column English, an optional free-text description/context column, then
one column per target language — matches the agency's existing glossary
doc format. Re-uploading replaces the whole project glossary.

A check for one language only ever needs EN + RU + that one target
column (terms_for_language) — never the whole 35-language table, so
every check stays small and cheap regardless of glossary size.
"""
import io

import openpyxl

from app.excel_multi import _find_header_row, _is_meta_col, _normalize_lang_label

DESCRIPTION_COL_NAMES = {
    "description", "context", "comment", "comments", "note", "notes",
    "explanation", "пояснение", "описание", "комментарий", "примечание",
}
EN_COL_NAMES = {"en", "english"}


def parse_glossary_workbook(file_bytes: bytes) -> list[dict]:
    wb = openpyxl.load_workbook(io.BytesIO(file_bytes), data_only=True)
    ws = wb[wb.sheetnames[0]]
    if ws.max_row < 2:
        return []

    header_row = _find_header_row(ws)
    en_col = None
    description_col = None
    lang_cols: dict[int, str] = {}

    for c in range(1, ws.max_column + 1):
        v = ws.cell(row=header_row, column=c).value
        label = v.strip() if isinstance(v, str) else ""
        low = label.lower()
        if low in EN_COL_NAMES or (en_col is None and c == 1):
            # The glossary doc always starts with the English term, even if
            # that first column's header is blank or oddly labeled.
            en_col = c
        elif low in DESCRIPTION_COL_NAMES:
            description_col = c
        elif not label or _is_meta_col(label):
            continue
        else:
            # Accepts both a plain code ("es-mx") and the display style
            # ("ES (MX)") the agency also uses — both resolve to "es-mx".
            normalized = _normalize_lang_label(label)
            if " " in normalized or len(normalized) > 12:
                continue
            lang_cols[c] = normalized

    if en_col is None:
        en_col = 1

    terms = []
    for r in range(header_row + 1, ws.max_row + 1):
        en_val = ws.cell(row=r, column=en_col).value
        en_text = str(en_val).strip() if en_val else ""
        if not en_text:
            continue

        desc_text = ""
        if description_col:
            raw = ws.cell(row=r, column=description_col).value
            desc_text = str(raw).strip() if raw else ""

        translations = {}
        for c, code in lang_cols.items():
            raw = ws.cell(row=r, column=c).value
            text = str(raw).strip() if raw else ""
            if text:
                translations[code] = text

        terms.append({"term_en": en_text, "description": desc_text, "translations": translations})

    return terms


def terms_for_language(all_terms: list[dict], lang_code: str) -> list[dict]:
    """Project the full glossary down to just EN + RU + one target language
    — the only slice any single check ever needs."""
    lang_code = lang_code.lower()
    rows = []
    for t in all_terms:
        target = t["term_en"] if lang_code == "en" else t["translations"].get(lang_code, "")
        ru = t["term_en"] if lang_code == "ru" else t["translations"].get("ru", "")
        if not target and not ru:
            continue
        rows.append({"en": t["term_en"], "description": t["description"], "ru": ru, "target": target})
    return rows


def format_glossary_prompt(rows: list[dict], lang_code: str) -> str:
    if not rows:
        return ""
    lines = [f"Глоссарий проекта (формат строки: EN | RU | {lang_code} | пояснение, «—» = нет перевода на этот язык):"]
    for r in rows:
        line = f'{r["en"]} | {r["ru"] or "—"} | {r["target"] or "—"}'
        if r["description"]:
            line += f' | {r["description"]}'
        lines.append(line)
    return "\n".join(lines)
