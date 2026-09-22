"""
Parsing and checking of "Мульти" Excel uploads — the format Crowdin exports:
one sheet (sometimes several) with a header row of language codes, a
"Context" column naming each string, an optional "Max. length" column, and
one row per translatable string.
"""
import asyncio
import io
import re

import openpyxl

from app.gemini_client import run_gemini_checks_batch
from app.claude_client import (
    REGISTER_MIXED_TYPE,
    REGISTER_VALUE_TYPE,
    _ai_failure_warning,
    _filter_findings_by_checks,
    _is_hard_language,
    _model_for_lang,
    _register_mixed_finding,
    _truncation_warning,
    _usage_cost,
    build_batch_prompt,
    build_register_report,
    cancel_message_batch,
    create_message_batch,
    get_batch_results,
    get_batch_status,
    group_batch_findings,
    parse_json_array,
    run_ai_checks_batch,
)
from app.rule_checks import run_rule_checks

LANG_CODE_RE = re.compile(r"^[a-z]{2,3}(-[a-z0-9]{2,5})?$")
AI_CONCURRENCY = 5

# Industry-standard placeholder: a translator (or the client) can mark a
# specific cell "DO NOT TRANSLATE" to mean this string is deliberately left
# as-is for this language on purpose (a brand name, a code, a string that's
# only needed in one of several languages) — not a missing or wrong
# translation. Александр's files use this in exactly that way: some rows
# need translating into every language, others explicitly don't for a given
# one. Recognized case-insensitively, with or without surrounding brackets.
_DO_NOT_TRANSLATE_RE = re.compile(r"^[\[\(]?\s*do\s+not\s+translate\s*[\]\)]?$", re.IGNORECASE)


def _is_do_not_translate(text: str) -> bool:
    return bool(_DO_NOT_TRANSLATE_RE.match((text or "").strip()))

# Above this much combined text (characters, summed across every checkable
# row × every target language — a rough proxy for total AI cost and how
# long a synchronous run would take), a multi-check is submitted through
# Anthropic's Message Batches API instead of run live: half the per-token
# price, but results typically land within an hour instead of a couple of
# minutes. A short single-language check, or a handful of languages, stays
# under this and runs instantly as before — but a small document checked
# across many languages (which is what actually drives the per-token cost
# up) will typically cross this threshold too. This is a starting guess,
# not a measured number — tune it based on real usage once there's a
# track record of actual costs and turnaround times.
BATCH_THRESHOLD_CHARS = 10_000

# How many (context, source, translation) triples go into ONE AI call, for
# either processing path. Александр noticed (2026-09-17) that a document
# check on a large language could report "всё чисто" while pasting the very
# same problem row into the single-pair fields DID find something — the
# most likely cause: a single AI call covering hundreds of rows of one
# language at once has to split its attention across all of them, and a
# finding that needs actual reading comprehension (a subtle mistranslation,
# a dropped grammatical particle) is exactly the kind that can get missed
# under that load, unlike a mechanical rule check (numbers, placeholders),
# which is never done by the AI at all and so never suffers from this.
# Splitting one language's rows into smaller chunks — each its own separate
# AI call — keeps that per-row attention high regardless of how big the
# whole document is. Kept deliberately small, per Александр's explicit
# choice, to prioritize per-row attentiveness over cost/latency.
#
# Trade-off this comes with: the "Повторяется по всему документу" repeat
# dedup (see BATCH_PROMPT/group_batch_findings/_resolve_repeated_findings)
# can only ever recognize a repeat within rows that land in the SAME chunk —
# the model literally never sees rows from a different chunk in the same
# call. A term mistranslated identically across, say, 60 rows of one
# language will now come back as a handful of "Повторяется..." findings
# (one per chunk it shows up in) rather than exactly one for the whole
# document. Still far better than one finding per row, just not perfect
# document-wide dedup any more.
MAX_ROWS_PER_AI_CALL = 15

# Александр's real-world test, 2026-09-22: the SAME Marathi pair ("फ्रибет
# без отыгрыша"), through the SAME model (Opus — mr is on HARD_LANGUAGE_BASES,
# see claude_client._is_hard_language), with the SAME prompt/instructions —
# was MISSED when checked as part of a normal batch, but CAUGHT when checked
# alone via the single-pair form. Confirms the exact mechanism
# MAX_ROWS_PER_AI_CALL above already exists to fight (splitting attention
# across many rows in one call), just that 15 rows is still too many for the
# hardest languages specifically — only ever tested/proven down to 1 row at
# a time, so that's what this uses rather than guessing an untested middle
# value like 3 or 5. Costs noticeably more per hard-language row (each call
# repeats the full instruction text for just one row instead of sharing it
# across up to 15) — a deliberate quality-over-cost trade-off, same spirit
# as MAX_ROWS_PER_AI_CALL's own. Easy/normal languages are unaffected —
# still MAX_ROWS_PER_AI_CALL as before.
MAX_ROWS_PER_AI_CALL_HARD = 1


def _chunk_size_for_lang(target_lang: str) -> int:
    return MAX_ROWS_PER_AI_CALL_HARD if _is_hard_language(target_lang) else MAX_ROWS_PER_AI_CALL


def _chunk_list(items: list, size: int) -> list[list]:
    """Splits items into consecutive chunks of at most `size`, preserving
    order — the empty list yields no chunks at all (not one empty chunk)."""
    return [items[i:i + size] for i in range(0, len(items), size)]

# Known non-language metadata column names seen in Crowdin-style exports.
# Anything NOT in this list and not Context/Max length is treated as a
# language column — real client files have non-standard codes (a 4-letter
# code, a look-alike Cyrillic character typo'd into a code, etc.) that a
# strict regex would silently drop, which is worse than being permissive.
META_COL_NAMES = {
    "context", "key", "id", "string id", "identifier", "comment",
    "status", "screenshot", "reference", "notes", "note",
    "контекст", "ключ", "тз", "комментарий", "примечание", "статус",
}

# A "label: number" shape — "NOTIF title: 20", "Лимиты: PUSH banner: 25" —
# is unambiguously a character-limit/spec column, never a language: a real
# language code never contains a colon. Matched at the end of the header so
# a "Лимиты: ..." prefix in front doesn't matter. Also treat any header that
# starts with the Russian word for "limits" (a merged section header like a
# bare "Лимиты" spanning several columns, with no colon of its own) the
# same way. Client files pack in columns like this alongside the real
# language columns — Александр flagged them as noise in the "not
# recognized as languages" notice, since it's obvious on sight (and by
# comparing with the source column) that they were never meant to be one.
_LIMIT_SPEC_RE = re.compile(r":\s*\d+\s*$")


def _is_context_col(header: str) -> bool:
    return header.strip().lower() == "context"


def _is_max_length_col(header: str) -> bool:
    h = header.strip().lower()
    return "max" in h and "length" in h


def _is_meta_col(header: str) -> bool:
    return header.strip().lower() in META_COL_NAMES


def _is_limit_spec_col(header: str) -> bool:
    h = header.strip()
    return bool(_LIMIT_SPEC_RE.search(h)) or h.lower().startswith("лимит")


_PAREN_LANG_RE = re.compile(r"^([a-zA-Zа-яА-Я]{2,3})\s*\(([a-zA-Zа-яА-Я0-9]{1,5})\)$")

# The same "ES (MX)"/"PT (BR)" display style, but without the parentheses —
# "ES MX", "PT BR" — a header written this way isn't currently reachable at
# all: with no hyphen and no parens for _PAREN_LANG_RE to key off, it falls
# straight through to a plain .lower() with the space still in it, which
# parse_workbook's own space check then rejects as "looks like prose, not a
# language code". Deliberately requires BOTH words to be fully uppercase in
# the source file (unlike the parenthesized form above, which is
# case-insensitive) — that's what tells a genuine "ES MX"/"PT BR"-style code
# apart from an ordinary two-word column header like "Task name" or a
# capitalized "No data", which are never written in ALL CAPS in Александр's
# real files. Deliberately does NOT try to be clever about the second
# word's length ("PT BR" and "ES MX" happen to be 2+2, but keep the same
# 1-5 character allowance as the parenthesized form for other real cases).
_SPACE_LANG_RE = re.compile(r"^([A-Z]{2,3})\s+([A-Z0-9]{1,5})$")


# Unlike Spanish (which the agency always spells out by region — es-ar,
# es-mx, es-es — because it genuinely handles several), Portuguese has only
# ever meant one thing in Александр's work: Brazilian Portuguese. But
# Portugal's and Brazil's Portuguese differ enough (vocabulary, formality
# conventions) that the AI check should be told explicitly which one it's
# dealing with, exactly the way an explicit "es-ar" already tells it
# Argentine Spanish rather than leaving that to be guessed from a bare "es"
# (see _target_lang_line in app.claude_client — whatever code reaches it as
# target_lang is what the AI is told the language IS). So a bare,
# unqualified "pt" — no region attached at all — defaults to Brazilian
# Portuguese ("pt-br") wherever a language code is first minted: a file
# column header, or a manually-typed catalog addition. An EXPLICIT region
# ("PT (PT)", "pt-pt") is left alone — if Portugal's own Portuguese is ever
# actually needed, spelling it out that way is how to ask for it instead of
# getting the Brazilian default.
_DEFAULT_REGION_FOR_BARE_LANG = {
    "pt": "pt-br",
}


def _normalize_lang_label(label: str) -> str:
    """Turns a display-style language header like "ES (MX)", "PT (BR)", or
    the same without parentheses ("ES MX", "PT BR" — must be ALL CAPS, see
    _SPACE_LANG_RE) into the hyphenated form used everywhere else
    ("es-mx", "pt-br"), while leaving an already-plain code like "es-mx" or
    "fr-сi" untouched. Some client files use one style, some another, so
    all of them need to resolve to the same lang_code. A handful of bare
    codes also get defaulted to a specific region here — see
    _DEFAULT_REGION_FOR_BARE_LANG — since an explicit region ("PT (BR)",
    "pt-pt") above already means something specific and must never be
    overridden by that default."""
    label = label.strip()
    m = _PAREN_LANG_RE.match(label)
    if m:
        return f"{m.group(1).lower()}-{m.group(2).lower()}"
    m = _SPACE_LANG_RE.match(label)
    if m:
        return f"{m.group(1).lower()}-{m.group(2).lower()}"
    code = label.lower()
    return _DEFAULT_REGION_FOR_BARE_LANG.get(code, code)


def _label_to_code(label: str, alias_map: dict[str, str] | None = None) -> str:
    """The single choke point for "what language does this raw text mean" —
    Александр's own idea: rather than every new nonstandard abbreviation
    (GEO, a mistaken "PR" for Portuguese, "HING" for Hinglish, whatever
    comes next) needing a code change from a developer, any manager can
    teach the platform a spelling once, through the global /language-
    aliases dictionary (see models.LanguageAlias), and it's recognized
    everywhere from then on — a file's own column header, the
    Tone-of-address document, or a manually-typed catalog addition.

    alias_map (built fresh from that table by the caller — this function
    stays a pure string transform, no DB access of its own, so it can
    still be unit-tested with a plain dict) is checked FIRST, against the
    raw label exactly as typed (trimmed, case-insensitive) — deliberately
    BEFORE _normalize_lang_label's own regex-based rules, and deliberately
    able to rescue a label that wouldn't otherwise even look language-
    shaped at all (contains a space, isn't 2-3 letters, whatever) since
    the whole point is to keep shrinking how often anything ends up
    unrecognized. A label taught this way always wins over a coincidental
    regex match. Falls back to the existing _normalize_lang_label rules
    when nothing in the dictionary matches."""
    hit = (alias_map or {}).get(label.strip().lower())
    if hit:
        return hit
    return _normalize_lang_label(label)


def _base_lang(code: str) -> str:
    """The base language subtag of a code — "ko" from "ko-KR", "es" from
    "es-mx", or the whole thing if it has no region part."""
    return code.strip().lower().split("-")[0]


def _subtags(code: str) -> set[str]:
    """Every hyphen-separated piece of a code, lowercased — {"kk", "kz"}
    for "kk-KZ". Used to match a bare code against either the language part
    OR the region part of a fuller one, since the agency's own documents
    are inconsistent about which they use for a short label — the Tone doc
    labels some languages by their country's code alone ("KZ" for Kazakh,
    "TJ" for Tajik, "BD" for Bengali) rather than the actual ISO language
    subtag ("kk", "tg", "bn") that a region-qualified code like "kk-KZ"/
    "tg-TJ"/"bn-BD" would use — a plain prefix-only match would miss these
    entirely, since "kz" isn't the base of "kk-kz"."""
    return set(code.strip().lower().split("-"))


# Bare codes that are real, independent ISO-639 languages in this project's
# own right — every one of them actually appears as its own language column
# somewhere in Александр's real files. Used by _language_subtags_compatible
# below to tell apart two very different reasons a short code might share
# letters with a region subtag: the agency's own country-code-style
# shorthand for a language (Tone doc's "KZ" for Kazakh, "TJ" for Tajik,
# "BD" for Bengali — none of which are themselves real ISO-639 codes, so
# matching them against a region subtag is exactly the intended trick), vs
# a genuine ISO-639 language code that PURELY BY COINCIDENCE also spells a
# real but unrelated country's ISO-3166 code (Arabic "ar" vs Argentina's
# country code "AR"; also latent landmines for the same reason even though
# no real file has hit them yet: Bengali "bn"/Brunei "BN", Kyrgyz "ky"/
# Cayman Islands "KY", Marathi "mr"/Mauritania "MR", Tajik "tg"/Togo "TG",
# Tagalog "tl"/Timor-Leste "TL"). A code in this set must never be treated
# as merely someone else's region fragment.
INDEPENDENT_LANGUAGE_CODES = {
    "ar", "az", "bn", "de", "el", "en", "es", "fr", "hi", "id", "it", "ja",
    "kk", "ko", "ky", "mr", "ms", "my", "pl", "pt", "ro", "ru", "sw", "te",
    "tg", "th", "tl", "tr", "uk", "ur", "uz", "vi", "zh",
}


def _language_subtags_compatible(a: str, b: str) -> bool:
    """True when two language codes plausibly name the same language once
    bridged across granularity — the shared logic behind both
    resolve_lang_code's and merge_lang_codes's "same language, differently
    spelled" bridging — while refusing a match that only "works" because
    an ISO-639 language code happens to spell the same two letters as an
    unrelated ISO-3166 country code (see INDEPENDENT_LANGUAGE_CODES above:
    Arabic "ar" must never match "es-ar" — Spanish, Argentina — just
    because Argentina's country code is also "AR").

    Always safe: the two codes share the same LANGUAGE subtag — "ko" and
    "ko-KR", or "es-ar" and "es-mx" by their common "es".

    Also safe, but only for a bare code that ISN'T itself a real,
    independent language (the agency's own country-code-style shorthand —
    "KZ" for Kazakh, "TJ" for Tajik, "BD" for Bengali, which aren't
    themselves recognized language codes) matching the REGION half of a
    fuller code ("KZ" against "kk-KZ"). A bare code that IS a real
    language in its own right is excluded from this side of the match
    entirely — it may only match by sharing an actual LANGUAGE subtag,
    never by coincidentally matching someone else's region."""
    a, b = a.strip().lower(), b.strip().lower()
    if _base_lang(a) == _base_lang(b):
        return True
    a_is_shorthand = "-" not in a and a not in INDEPENDENT_LANGUAGE_CODES
    b_is_shorthand = "-" not in b and b not in INDEPENDENT_LANGUAGE_CODES
    if a_is_shorthand and a in _subtags(b):
        return True
    if b_is_shorthand and b in _subtags(a):
        return True
    return False


def _freeze(value):
    """Makes a value hashable/comparable for the equality check in
    resolve_lang_code below — a document's row might be a plain string
    (e.g. Tone's register) or a dict of fields, depending on the doc."""
    if isinstance(value, dict):
        return tuple(sorted(value.items()))
    return value


def resolve_lang_code(requested: str, available, values: dict | None = None):
    """Matches a requested language code against a set/dict/iterable of
    codes actually present in one document, bridging granularity mismatches
    within that document — e.g. a target language selected as a plain "ko"
    still finds a region-qualified "ko-KR" row, and vice versa.

    Tries an exact (case-insensitive) match first. Failing that, falls
    back to matching the requested code against a compatible subtag of an
    available code (see _language_subtags_compatible — language part
    always, region part only for a genuine country-code-style shorthand,
    never for a bare code that's a real language in its own right) — but
    ONLY when exactly one available code is compatible; if several are
    (e.g. a document has both "es-ES" and "es-AR" with genuinely different
    rules for each), guessing would silently apply the wrong regional
    rule, so this returns None instead — the caller then treats the
    language as if it had no entry at all, exactly like today's "document
    uploaded but this language is missing" case, rather than picking one
    region at random.

    Some languages genuinely split into a handful of regional variants
    where every OTHER variant besides one or two special cases shares the
    same rule (e.g. Spanish: Spain and Argentina each have their own
    format, but "the rest of Latin America" is one shared format that can
    show up under any of several country codes — es-MX, es-CL, es-PE...).
    The agency already has a way to express that directly in the document
    itself: list every code that shares one row together in one cell,
    separated by "/" (already used for French: "fr-CI / fr-FR") — every
    code listed then resolves by exact match, no ambiguity at all. As a
    safety net for a code that WASN'T listed, if `values` (a {code: value}
    mapping) is passed, several same-base candidates still resolve when
    they all happen to carry the identical value — applying it is safe
    regardless of which one is picked, since they don't actually disagree.

    Returns the matching code from `available` (preserving its original
    casing), or None if nothing resolves safely.
    """
    requested = (requested or "").strip().lower()
    if not requested:
        return None

    avail_list = list(available)
    by_lower = {a.lower(): a for a in avail_list}
    if requested in by_lower:
        return by_lower[requested]

    matches = [a for a in avail_list if _language_subtags_compatible(requested, a)]
    if len(matches) == 1:
        return matches[0]
    if len(matches) > 1 and values is not None:
        distinct = {_freeze(values[a]) for a in matches if a in values}
        if len(distinct) == 1:
            return matches[0]
    return None


def _lang_selected(lang: str, target_langs_filter: set[str]) -> bool:
    """Whether `lang` — a code straight from a FILE's own header, already
    normalized by _normalize_lang_label — should count as "selected" by a
    manager's target_langs_filter (the raw codes of whichever catalog
    checkboxes were ticked in the UI). A plain `in` check breaks the
    instant the file's own spelling doesn't EXACTLY match the catalog's
    spelling for the same language: a catalog still holding an older bare
    "pt" entry (from before Portuguese started defaulting to "pt-br") next
    to a freshly-uploaded file whose "PT" column now normalizes to
    "pt-br" is exactly the case that motivated this — Александр's
    detect-languages notice already told him "pt-br" would be checked
    (that endpoint has always bridged catalog/file spelling mismatches),
    so the actual check run must honor that instead of silently dropping
    the language over a spelling technicality.

    Deliberately narrower than resolve_lang_code itself: bridges ONLY when
    exactly one of the two sides is a bare code (no region at all) and the
    other is region-qualified — bare "pt" <-> file's "pt-br", or a
    country-code-style shorthand like "kz" <-> "kk-KZ" — via the same
    _language_subtags_compatible rules used everywhere else. Two codes
    that are BOTH already region-qualified are never bridged just because
    they happen to share a base language: "es-mx" and "es-es" are
    deliberately different, explicitly-added catalog languages (see
    merge_lang_codes), and ticking one must never silently sweep in the
    other's column too — resolve_lang_code's own base-language shortcut is
    too permissive for that case, so it isn't reused here. Refuses (rather
    than guessing) when a single FILE column could bridge to more than one
    ticked filter entry — e.g. a file's bare "es" column with both "es-ar"
    and "es-mx" ticked never silently ends up checked as just one of them.

    This is NOT symmetric with resolve_lang_code's own ambiguity refusal,
    and deliberately so: when the CATALOG side is the bare one instead
    (only a generic "es" ticked) and the file has several explicit
    columns for it ("es-ar" AND "es-mx" both present), each column bridges
    to that one bare entry independently and BOTH get selected — the
    manager's bare tick reads as "check Spanish, generically", and this
    module's guiding rule (see the PR/Peru and GEO catalog fixes elsewhere
    in this file) is that a language never silently drops out of a check
    just because of a granularity mismatch. Over-including a language the
    manager arguably meant to cover is a far smaller problem than the bug
    this function exists to fix (a language silently never checked at
    all), so this asymmetry is intentional, not a gap to close."""
    lang_low = lang.strip().lower()
    filt = {f.strip().lower() for f in target_langs_filter if f and f.strip()}
    if lang_low in filt:
        return True
    lang_is_bare = "-" not in lang_low
    matches = [
        f for f in filt
        if (("-" not in f) != lang_is_bare) and _language_subtags_compatible(lang_low, f)
    ]
    return len(matches) == 1


def merge_lang_codes(codes) -> list[str]:
    """Deduplicates a set of language codes gathered from several documents
    for display (e.g. the target-language picker) — when a base language
    has only one region variant across everything ("ko" in one doc, "ko-KR"
    in another), it's shown once, as its more specific spelling. When a
    base language genuinely has several distinct region variants (es-ES,
    es-AR, es-MX), each is kept as its own separate entry, since they mean
    different formatting rules and must be picked explicitly.

    Region-qualified codes are grouped by their first (language) subtag
    only — deliberately narrower than a full compatibility check, since
    two region-qualified codes should never merge just for sharing a
    region (hi-IN and mr-IN are different languages that happen to both be
    spoken in India). A BARE code with no region of its own is looser by
    nature — it's merged into whichever region-qualified group it's
    compatible with (see _language_subtags_compatible: a genuine
    country-code-style shorthand like Tone's "KZ" for Kazakh matches on
    ANY subtag, but a bare code that's a real independent language, like
    Arabic "ar", only matches by sharing an actual language subtag — it
    must never be absorbed into another language's group just because it
    happens to spell the same two letters as one of that group's REGIONS,
    e.g. Arabic "ar" vs Argentina's country code inside "es-ar"), and only
    when that's unambiguous (exactly one group matches)."""
    codes = [(c or "").strip() for c in codes if c and c.strip()]
    hyphenated = [c for c in codes if "-" in c]
    bare = [c for c in codes if "-" not in c]

    groups: dict[str, list[str]] = {}
    for code in hyphenated:
        groups.setdefault(_base_lang(code), []).append(code)

    for code in bare:
        low = code.lower()
        matching_bases = {
            base for base, variants in groups.items()
            if any(_language_subtags_compatible(low, v) for v in variants)
        }
        if len(matching_bases) == 1:
            groups[next(iter(matching_bases))].append(code)
        else:
            groups.setdefault(low, []).append(code)

    result = []
    for base, variants in groups.items():
        distinct = sorted(set(variants))
        regioned = [v for v in distinct if "-" in v]
        if len(regioned) >= 2:
            # Several genuinely different regional variants (es-ES vs
            # es-AR vs es-MX) — keep each; drop any bare/generic spelling
            # of the same base, since it's ambiguous which region it means.
            result.extend(regioned)
        elif len(regioned) == 1:
            # Only one spelling actually matters for this language — a
            # bare "ko" alongside it is the same language, not a second one.
            result.append(regioned[0])
        else:
            result.append(distinct[0])
    return sorted(result)


def _find_header_row(ws, max_scan: int = 5) -> int:
    best_row, best_score = 1, -1
    for r in range(1, min(max_scan, ws.max_row) + 1):
        score = 0
        for c in range(1, ws.max_column + 1):
            v = ws.cell(row=r, column=c).value
            if isinstance(v, str) and LANG_CODE_RE.match(v.strip().lower()):
                score += 1
        if score > best_score:
            best_row, best_score = r, score
    return best_row


def parse_workbook(file_bytes: bytes, alias_map: dict[str, str] | None = None) -> list[dict]:
    """Returns a list of parsed sheets: each with lang codes found and rows.

    alias_map: the manager-built global "this raw spelling means this
    language" dictionary (see _label_to_code and models.LanguageAlias),
    fetched fresh from the DB by the caller. Checked before a column is
    ever judged "not language-shaped" — a taught alias can rescue a
    header that the regex-only rules would otherwise drop into
    `unrecognized_columns`, which is the whole point of the dictionary
    existing at all."""
    wb = openpyxl.load_workbook(io.BytesIO(file_bytes), data_only=True)
    sheets = []

    for sheet_name in wb.sheetnames:
        ws = wb[sheet_name]
        if ws.max_row < 2:
            continue

        header_row = _find_header_row(ws)
        context_col = max_length_col = None
        lang_cols: dict[int, str] = {}
        unrecognized: list[str] = []

        for c in range(1, ws.max_column + 1):
            v = ws.cell(row=header_row, column=c).value
            if not isinstance(v, str) or not v.strip():
                continue
            label = v.strip()
            alias_hit = (alias_map or {}).get(label.lower())
            if alias_hit:
                # An explicitly taught spelling always wins, whatever it
                # looks like — checked before ANY other classification
                # (context/max-length/meta/limit-spec, and the "looks like
                # prose" guess below), so a taught alias can never be
                # silently swallowed by one of those shortcuts. This is
                # what "checked before a column is ever judged 'not
                # language-shaped'" (see the docstring) actually means —
                # it used to only run after the meta/limit-spec `continue`s,
                # which could drop a taught alias whose raw text happened
                # to also match one of those unrelated shapes.
                lang_cols[c] = alias_hit
                continue
            if _is_context_col(label):
                context_col = c
            elif _is_max_length_col(label):
                max_length_col = c
            elif _is_meta_col(label):
                continue
            elif _is_limit_spec_col(label):
                # Obviously a limit/spec column, not a near-miss language
                # code — drop it silently instead of flagging it to the
                # manager as "not recognized as a language".
                continue
            else:
                normalized = _normalize_lang_label(label)
                if " " in normalized or len(normalized) > 12:
                    # Looks like prose, not a language code — skip rather
                    # than misread a stray comment column as a "language".
                    unrecognized.append(label)
                else:
                    lang_cols[c] = normalized

        if not lang_cols:
            continue

        # A language code assigned to 2+ columns (e.g. two columns both
        # headed "ru") is a real structural ambiguity, not just cosmetic —
        # the per-row `values` dict comprehension just below can only keep
        # ONE column's value per code, and silently lets the rightmost
        # column win with no record anywhere that the other one was
        # dropped. Recorded here (column order preserved) so callers can
        # warn the manager about it instead of leaving it invisible —
        # caught live on Александр's real file (2026-09-22): a duplicated
        # "ru" header at both C1 and AA1, which doubled his "missing
        # translation" findings for ru and quietly discarded whichever
        # column wasn't C1/rightmost — see _duplicate_language_warning.
        col_letters_by_code: dict[str, list[str]] = {}
        for c, code in lang_cols.items():
            col_letters_by_code.setdefault(code, []).append(openpyxl.utils.get_column_letter(c))
        duplicate_language_columns = {
            code: letters for code, letters in col_letters_by_code.items() if len(letters) > 1
        }

        rows = []
        last_context = ""
        for r in range(header_row + 1, ws.max_row + 1):
            # When a code maps to 2+ columns (see duplicate_language_columns
            # above), a plain {code: ...} dict comprehension over lang_cols
            # would just let the LAST column blindly overwrite every earlier
            # one, blank or not — and that's exactly how a real, actively
            # used bug hid for this long: Александр's real file has a
            # completely empty second "ru" column sitting to the right of
            # the one with his actual Russian text, so "rightmost wins"
            # silently produced an EMPTY source for every single row
            # whenever "ru" was picked as the source language — this is the
            # true cause of his very first "Источник пусто" report from
            # earlier in this project, not merged cells as first suspected.
            # A blank duplicate column is essentially never the intended
            # one, so this instead keeps the last NON-BLANK value seen,
            # falling back to blank only if every duplicate for that code is
            # genuinely blank on this row. Columns that actually DISAGREE
            # (two different real values) still resolve to the rightmost —
            # unchanged from before, and still flagged to the manager via
            # duplicate_language_columns/the warning it drives, since THAT
            # kind of collision genuinely needs a human decision.
            values: dict[str, object] = {}
            for c, code in lang_cols.items():
                v = ws.cell(row=r, column=c).value
                is_blank = v is None or str(v).strip() == ""
                if code not in values or not is_blank:
                    values[code] = v
            if all(v is None or str(v).strip() == "" for v in values.values()):
                continue

            context = ""
            if context_col:
                raw_ctx = ws.cell(row=r, column=context_col).value
                context = str(raw_ctx).strip() if raw_ctx else ""
            if context:
                last_context = context
            else:
                context = last_context

            max_length = None
            if max_length_col:
                raw_ml = ws.cell(row=r, column=max_length_col).value
                if isinstance(raw_ml, (int, float)):
                    max_length = int(raw_ml)

            rows.append({
                "excel_row": r,
                "context": context,
                "max_length": max_length,
                "values": {code: ("" if v is None else str(v)) for code, v in values.items()},
            })

        sheets.append({
            "sheet_name": sheet_name,
            # Deduplicated — a code appearing in 2+ columns must not appear
            # twice in this list. Before this fix it did, and every caller
            # that builds `target_langs` from it (run_multi_check,
            # estimate_check_volume, build_batch_plan) would check that one
            # language TWICE — doubling both its AI cost and its findings —
            # on top of the ambiguity `duplicate_language_columns` itself
            # already warns about.
            "languages": sorted(set(lang_cols.values())),
            "rows": rows,
            "unrecognized_columns": unrecognized,
            "duplicate_language_columns": duplicate_language_columns,
        })

    return sheets


def pick_source_lang(sheets: list[dict], preferred: str | None) -> str:
    """Resolves the manager's chosen source language (e.g. "ru", from the
    RU/EN buttons in the UI) against the language codes actually found as
    column headers in the uploaded file. Uses resolve_lang_code rather than
    a literal match, since the file's own column can be a differently
    granular spelling of the same language (e.g. "ru-RU" for a plain "ru")
    — Александр hit this: his file's Russian column wasn't spelled exactly
    "ru", the literal check silently missed it, and the source language
    silently fell back to English (whatever column happened to be labeled
    "en-001"), even though he'd picked Russian. Falls back to English, then
    alphabetically first, only when the requested language truly isn't in
    the file at all."""
    all_langs: set[str] = set()
    for s in sheets:
        all_langs.update(s["languages"])
    if preferred:
        resolved = resolve_lang_code(preferred, all_langs)
        if resolved:
            return resolved
    if "en" in all_langs:
        return "en"
    return sorted(all_langs)[0] if all_langs else "en"


def _duplicate_language_warning(code: str, columns: list[str], is_source: bool) -> dict:
    """Same synthetic-finding pattern as _truncation_warning/
    _ai_failure_warning (a "type": "system" entry, already rendered by the
    frontend with its own "⚠ Внимание" badge — no UI change needed) —
    makes an otherwise silent multi-column ambiguity visible instead of
    just looking like ordinary clean data.

    `columns` is in left-to-right sheet order (see parse_workbook). Which
    column's value actually gets used is now decided PER ROW there — a
    blank duplicate never wins over a non-blank one, only a genuine
    disagreement between two non-blank values falls back to "rightmost
    wins" — so this message can no longer name a single fixed "used"
    column the way it used to; it describes the rule instead. That
    per-row resolution is exactly what fixed Александр's real "Источник
    пусто" bug (2026-09): a totally empty second "ru" column, to the
    right of the one with his actual Russian text, used to silently win
    every row just for being rightmost."""
    if is_source:
        intro = (
            f"Внимание: в файле несколько столбцов помечены как исходный язык «{code}» "
            f"(столбцы {', '.join(columns)}). Это влияет на проверку СРАЗУ ВСЕХ языков перевода в этом "
            f"листе — показано здесь один раз. "
        )
    else:
        intro = (
            f"В файле несколько столбцов с одинаковым языковым кодом «{code}» "
            f"(столбцы {', '.join(columns)}). "
        )
    return {
        "type": "system",
        "severity": "high",
        "message": (
            intro
            + f"Платформа не может определить, какой из них правильный. Для каждой строки используется "
              f"непустое значение (если пустая только одна из колонок — берётся та, где есть текст), а "
              f"если заполнены обе и текст в них отличается — используется самая правая колонка, "
              f"{columns[-1]}. Проверьте, пожалуйста, структуру файла — возможно, один из этих столбцов "
              f"лишний или назван неправильно."
        ),
    }


# KNOWN GAP (low severity, caught by subagent review): both callers
# (run_multi_check, finalize_batch_results) only ever invoke the function
# below once per language actually being checked. If the manager's
# target_langs_filter ends up excluding every target language (or the file
# genuinely has none besides the duplicated source), there's no language
# "slot" left to attach the source's warning to, and it's silently absent
# from the check RESULT — though nothing was actually checked or billed in
# that state either, and the pre-check /multi-check/detect-languages screen
# (see app.main) already shows this exact ambiguity before "start"
# regardless of which target languages end up selected. Not worth
# restructuring the per-language result shape to cover a case where
# nothing is being checked at all.
def _apply_duplicate_language_warnings(
    findings_list: list[dict],
    lang: str,
    duplicate_language_columns: dict[str, list[str]],
    source_lang: str,
    show_source_warning: bool,
) -> list[dict]:
    """Appends a synthetic warning "row" (excel_row=0, same pattern as the
    truncation/AI-failure warnings) for any duplicate-header ambiguity that
    affects THIS language's own check — its own column is duplicated, or
    (once per sheet, via `show_source_warning`) the shared source column
    is duplicated, which affects every language equally."""
    warnings = []
    if lang in duplicate_language_columns:
        warnings.append(_duplicate_language_warning(lang, duplicate_language_columns[lang], is_source=False))
    if show_source_warning and source_lang in duplicate_language_columns:
        warnings.append(
            _duplicate_language_warning(source_lang, duplicate_language_columns[source_lang], is_source=True)
        )
    if not warnings:
        return findings_list
    return findings_list + [{
        "excel_row": 0,
        "context": "⚠ Системное предупреждение",
        "source": "",
        "translation": "",
        "findings": warnings,
    }]


def _extract_register_values(grouped: dict[int, list[dict]]) -> tuple[dict[int, list[dict]], dict[int, str]]:
    """Pulls REGISTER_VALUE_TYPE entries (see app.claude_client) out of a
    {item index: [finding, ...]} dict, returning (the same dict with those
    entries removed, {index: value}). Used by both the live path
    (_check_language_for_sheet below) and the Message-Batches finalize
    path (finalize_batch_results further down) — both end up with the
    model's raw response in exactly this shape, just reached differently
    (an already-awaited call here vs. a polled batch result there).

    Must run BEFORE either caller's "if findings: show this row" check —
    a register_value entry exists for every checked row regardless of
    whether there's a real problem, so leaving it in would make every
    single row look like it has a finding.

    A "mixed" value — this ONE row's own translation switches between
    «ты» and «вы» within itself, rather than using one consistently
    (Александр's ask, 2026-09-17: a single Excel cell can hold several
    sentences/paragraphs, and the tone can genuinely drift mid-cell) — is
    deliberately NOT put into the returned values dict at all: it's not a
    vote for the document's majority tone or a counted exception, it's a
    real problem on this specific row, so it's turned into an ordinary
    visible finding right here instead (via _register_mixed_finding) and
    kept in the row's own findings list, same as any other real finding."""
    cleaned: dict[int, list[dict]] = {}
    values: dict[int, str] = {}
    for idx, findings in grouped.items():
        kept = []
        for f in findings:
            if f.get("type") == REGISTER_VALUE_TYPE:
                v = f.get("value")
                if v == "mixed":
                    kept.append(_register_mixed_finding())
                elif v in ("formal", "informal", "neutral"):
                    values[idx] = v
            else:
                kept.append(f)
        if kept:
            cleaned[idx] = kept
    return cleaned, values


def _resolve_repeated_findings(grouped: dict[int, list[dict]], rows: list[dict]) -> dict[int, list[dict]]:
    """A finding that named several pairs at once via "rows" in the model's
    response (see BATCH_PROMPT — Александр's ask, 2026-09-17: the same
    exact problem repeated identically across many rows shouldn't be N
    separate, near-duplicate findings) survives group_batch_findings as
    ONE finding attached to its first row, carrying the OTHER rows' item
    indices in an internal "_also_idx" key (item indices mean nothing
    outside this module). This resolves that into the actual Excel row
    numbers the manager sees in the report, folds them into the finding's
    own message, and strips the internal key so it never reaches the
    response. Runs on whatever ai_grouped looks like BEFORE the "if
    findings: show this row" step, same as _extract_register_values."""
    resolved: dict[int, list[dict]] = {}
    for idx, findings in grouped.items():
        new_findings = []
        for f in findings:
            also_idx = f.get("_also_idx")
            if also_idx:
                f = {k: v for k, v in f.items() if k != "_also_idx"}
                own_excel_row = rows[idx]["excel_row"] if 0 <= idx < len(rows) else None
                also_rows = sorted(
                    {rows[i]["excel_row"] for i in also_idx if 0 <= i < len(rows)}
                    - {own_excel_row}
                )
                if also_rows:
                    rows_str = ", ".join(str(r) for r in also_rows)
                    f["message"] = f"{f.get('message', '')} (также в строках: {rows_str})"
            new_findings.append(f)
        resolved[idx] = new_findings
    return resolved


def _count_real_findings(findings_list: list[dict]) -> int:
    """The "N проблем"/"N найдено" count shown across the UI (multi-check
    headline, per-language row counts, history list) — every real finding,
    EXCLUDING the synthetic register_summary report appended by
    _register_summary_block below. That report is a factual "here's the
    tone actually used" note, not a problem to fix, so a check that only
    ran "register" on an otherwise clean document must report 0 problems,
    not 1 per language — counting it here would contradict the whole point
    of dropping the old pass/fail tone-of-address judgment. Truncation/
    AI-failure warnings (type "system") are deliberately still counted —
    those genuinely are something the manager needs to notice.

    Also excludes calibration_debug findings (see
    _mark_calibration_debug_findings) — those are a separate, opt-in test
    signal shown alongside the real result, not part of it; counting them
    here would make turning on "🔬 Тест калибровки" alone inflate "N
    проблем" and look like the file got worse, which it didn't."""
    return sum(
        1
        for row in findings_list
        for f in row["findings"]
        if f.get("type") != "register_summary" and not f.get("calibration_debug")
    )


def _register_summary_block(report: dict | None) -> dict | None:
    """The synthetic "row" a per-language register report rides in as —
    same pattern already used for _truncation_warning/_ai_failure_warning
    (excel_row=0, a recognizable pseudo-context instead of a real row).
    None when there's nothing to report — register wasn't selected, no
    register_value entries came back at all, or none of the ones that did
    were classifiable as formal/informal (build_register_report itself
    returns None for that last case now too, 2026-09-18 — no more "не
    удалось определить" placeholder finding).

    report is build_register_report's structured return value — its
    "majority"/"exceptions"/"exception_labels" keys ride along on the
    finding itself (register_majority/register_exceptions/
    register_exception_labels) so the frontend can colorize «вы»/«ты» and
    highlight each exception's actual text (Александр's ask, 2026-09-17)
    without having to re-parse the plain-text message. "message" is still
    the plain-text fallback for the Excel export or any other plain-text-
    only reader — shortened 2026-09-18 to "Тон: <Вы|ты>[, кроме: ...]",
    down from a full "Тон обращения: везде на «вы»..." sentence."""
    if report is None:
        return None
    finding = {
        "type": "register_summary",
        "severity": "low",
        "message": f"Тон: {report['text']}.",
        "register_majority": report["majority"],
    }
    if report.get("exceptions") is not None:
        finding["register_exceptions"] = report["exceptions"]
    if report.get("exception_labels") is not None:
        finding["register_exception_labels"] = report["exception_labels"]
    return {
        "excel_row": 0,
        "context": "ℹ️ Тон обращения",
        "source": "",
        "translation": "",
        "findings": [finding],
    }


async def _run_ai_chunks(
    item_chunks: list[list[dict]],
    checks: list[str],
    extra_instructions: str,
    lang: str,
    source_lang: str,
    semaphore: asyncio.Semaphore,
    relaxed: bool = False,
) -> tuple[dict[int, list[dict]], float, bool]:
    """Runs every chunk of one language's items through the AI (bounded by
    the shared semaphore) and merges the per-chunk results back into a
    single {item index: findings} dict, with "_also_idx" indices shifted to
    match. Factored out of _check_language_for_sheet so the calibration
    debug pass (relaxed=True — see claude_client._calibration) can reuse
    the exact same chunking/merging logic as the normal, production pass
    instead of a second hand-rolled copy of it."""
    async def _run_chunk(chunk_items: list[dict]) -> tuple[dict[int, list[dict]], float, bool]:
        async with semaphore:
            return await run_ai_checks_batch(
                chunk_items, checks, extra_instructions, lang, source_lang, relaxed=relaxed,
            )

    chunk_results = await asyncio.gather(*[_run_chunk(c) for c in item_chunks])

    ai_findings_by_idx: dict[int, list[dict]] = {}
    cost_usd = 0.0
    truncated = False
    offset = 0
    for (chunk_findings, chunk_cost, chunk_truncated), chunk_items in zip(chunk_results, item_chunks):
        for local_idx, findings in chunk_findings.items():
            # "_also_idx" (see group_batch_findings) still holds indices
            # local to THIS chunk at this point — shift those too, or
            # _resolve_repeated_findings below would resolve them against
            # the wrong rows entirely.
            remapped = []
            for f in findings:
                if "_also_idx" in f:
                    f = {**f, "_also_idx": [offset + i for i in f["_also_idx"]]}
                remapped.append(f)
            ai_findings_by_idx[offset + local_idx] = remapped
        cost_usd += chunk_cost
        truncated = truncated or chunk_truncated
        offset += len(chunk_items)

    return ai_findings_by_idx, cost_usd, truncated


def _mark_calibration_debug_findings(findings: list[dict]) -> list[dict]:
    """Tags each finding from the relaxed-calibration debug pass (see
    calibration_debug below) so it's visibly distinct from a normal,
    production finding wherever findings get rendered — both a
    machine-readable "calibration_debug": True field (for the frontend, if
    it wants to style these differently) and a plain-text prefix on the
    message itself (so it's unmistakable even in the .xlsx report export,
    which just prints "message" as-is)."""
    out = []
    for f in findings:
        f = dict(f)
        f["calibration_debug"] = True
        f["message"] = "🔬 [Тест: мягкая калибровка, не обычный результат] " + str(f.get("message", ""))
        out.append(f)
    return out


async def _run_gemini_chunks(
    item_chunks: list[list[dict]],
    checks: list[str],
    extra_instructions: str,
    lang: str,
    source_lang: str,
    semaphore: asyncio.Semaphore,
) -> tuple[dict[int, list[dict]], float, bool, bool, str | None]:
    """Same chunking/merging/"_also_idx"-shifting as _run_ai_chunks above,
    but against Gemini (see app.gemini_client.run_gemini_checks_batch) —
    kept as its own function rather than a shared one because Gemini's
    per-chunk call returns two extra elements (errored, error_detail) that
    Claude's doesn't. Returns (findings keyed by item index, cost_usd,
    whether ANY chunk was truncated, whether ANY chunk failed to get a
    usable answer at all, a short reason why — from whichever failed chunk
    hit the problem first) — a language split across several chunks
    (MAX_ROWS_PER_AI_CALL) must still surface a warning even if only one of
    several chunks had a problem, not have it silently swallowed by the
    others that succeeded."""
    async def _run_chunk(chunk_items: list[dict]) -> tuple[dict[int, list[dict]], float, bool, bool, str | None]:
        async with semaphore:
            return await run_gemini_checks_batch(chunk_items, checks, extra_instructions, lang, source_lang)

    chunk_results = await asyncio.gather(*[_run_chunk(c) for c in item_chunks])

    findings_by_idx: dict[int, list[dict]] = {}
    cost_usd = 0.0
    truncated = False
    errored = False
    error_detail = None
    offset = 0
    for (chunk_findings, chunk_cost, chunk_truncated, chunk_errored, chunk_error_detail), chunk_items in zip(
        chunk_results, item_chunks
    ):
        for local_idx, findings in chunk_findings.items():
            remapped = []
            for f in findings:
                if "_also_idx" in f:
                    f = {**f, "_also_idx": [offset + i for i in f["_also_idx"]]}
                remapped.append(f)
            findings_by_idx[offset + local_idx] = remapped
        cost_usd += chunk_cost
        truncated = truncated or chunk_truncated
        errored = errored or chunk_errored
        error_detail = error_detail or chunk_error_detail
        offset += len(chunk_items)

    return findings_by_idx, cost_usd, truncated, errored, error_detail


def _mark_gemini_findings(findings: list[dict]) -> list[dict]:
    """Tags each finding from the optional "🌐 Проверить также через
    Gemini" pass (see gemini_check below) — a machine-readable
    "gemini_check": True field plus a "🌐 [Gemini] " message prefix,
    mirroring _mark_calibration_debug_findings's pattern. Unlike a
    calibration_debug finding, these ARE meant to be real, actionable
    findings — Александр explicitly chose to keep them counted in the
    headline "N проблем" (see run_multi_check/_count_real_findings, which
    only excludes calibration_debug and register_summary, not this). The
    tag says WHICH engine found it, not "ignore this, it's just a test"."""
    out = []
    for f in findings:
        f = dict(f)
        f["gemini_check"] = True
        f["message"] = "🌐 [Gemini] " + str(f.get("message", ""))
        out.append(f)
    return out


async def _check_language_for_sheet(
    sheet: dict,
    lang: str,
    source_lang: str,
    checks: list[str],
    extra_instructions: str,
    semaphore: asyncio.Semaphore,
    calibration_debug: bool = False,
    gemini_check: bool = False,
) -> tuple[list[dict], float, float, float]:
    """Returns (rows-with-findings, production cost_usd, extra cost_usd
    spent on the calibration_debug pass, extra cost_usd spent on the
    gemini_check pass — the last two always 0.0 when that feature is off)."""
    relevant_rows = []
    ai_items = []
    for row in sheet["rows"]:
        src = row["values"].get(source_lang, "")
        tgt = row["values"].get(lang, "")
        if not src.strip() and not tgt.strip():
            continue
        if _is_do_not_translate(tgt):
            continue
        relevant_rows.append(row)
        ai_items.append({"context": row["context"], "source": src, "translation": tgt})

    if not relevant_rows:
        return [], 0.0, 0.0, 0.0

    # See MAX_ROWS_PER_AI_CALL above — one big AI call covering the whole
    # language is split into several smaller ones instead, each still
    # bounded by the same shared semaphore (so the total number of AI calls
    # in flight at once across the whole upload is unaffected by chunking,
    # only how many rows any single call has to hold at once). A hard
    # language (see MAX_ROWS_PER_AI_CALL_HARD) gets an even smaller chunk
    # size — proven necessary, not just theoretical, by a real Marathi miss
    # that a single-pair check caught but a same-model, same-prompt batched
    # check didn't.
    item_chunks = _chunk_list(ai_items, _chunk_size_for_lang(lang))

    ai_findings_by_idx, cost_usd, truncated = await _run_ai_chunks(
        item_chunks, checks, extra_instructions, lang, source_lang, semaphore, relaxed=False,
    )
    ai_findings_by_idx = _resolve_repeated_findings(ai_findings_by_idx, relevant_rows)
    ai_findings_by_idx, register_values_by_idx = _extract_register_values(ai_findings_by_idx)

    # Александр's ask, 2026-09-22: find out whether real-world misses (found
    # by GPT/Gemini, not by us) come from the model's own knowledge gap or
    # from CALIBRATION_BASE's "only if confident" bar filtering out a
    # correct-but-uncertain finding — by running the SAME model, on the SAME
    # rows, a second time with only that one confidence bar loosened (see
    # claude_client.CALIBRATION_RELAXED_OPENING), and showing whatever it
    # additionally catches right alongside the normal result instead of
    # guessing. Deliberately NOT scoped to "hard" languages only — the
    # French {{country}} miss he wants covered too is on the ordinary model.
    debug_cost_usd = 0.0
    ai_findings_by_idx_relaxed: dict[int, list[dict]] = {}
    debug_truncated = False
    if calibration_debug:
        ai_findings_by_idx_relaxed, debug_cost_usd, debug_truncated = await _run_ai_chunks(
            item_chunks, checks, extra_instructions, lang, source_lang, semaphore, relaxed=True,
        )
        ai_findings_by_idx_relaxed = _resolve_repeated_findings(ai_findings_by_idx_relaxed, relevant_rows)
        ai_findings_by_idx_relaxed, _ = _extract_register_values(ai_findings_by_idx_relaxed)
        # _extract_register_values already strips raw REGISTER_VALUE_TYPE
        # entries, but a "mixed" value is turned INTO a visible
        # REGISTER_MIXED_TYPE finding right inside that same function (see
        # its own docstring) — register reporting isn't confidence-gated by
        # CALIBRATION_BASE at all (it's "report what's there", not "only if
        # sure"), so it has no business in a calibration comparison. Left
        # in, it would either falsely look like "the relaxed pass caught an
        # extra problem" (subagent review, 2026-09-22) or show as a
        # confusing duplicate 🔬 copy right next to the strict pass's own,
        # identical register_mixed finding for the same row.
        ai_findings_by_idx_relaxed = {
            idx: [f for f in findings if f.get("type") != REGISTER_MIXED_TYPE]
            for idx, findings in ai_findings_by_idx_relaxed.items()
        }
        ai_findings_by_idx_relaxed = {idx: fs for idx, fs in ai_findings_by_idx_relaxed.items() if fs}

    # "🌐 Проверить также через Gemini" — Александр's ask, 2026-09-22, after
    # a blind test (same prompt, no hints) showed Gemini independently
    # caught a real Marathi meaning error that Opus missed even with the
    # confidence bar loosened above, while correctly staying silent on a
    # genuinely-fine control example. Runs for EVERY checked language when
    # on (his own choice — not scoped to "hard" languages only), and its
    # findings are shown as real, actionable findings (not a debug-only
    # curiosity like calibration_debug above) — just clearly tagged with
    # which engine found them, since trust in a second provider is still
    # being built.
    gemini_cost_usd = 0.0
    gemini_findings_by_idx: dict[int, list[dict]] = {}
    gemini_truncated = False
    gemini_errored = False
    gemini_error_detail = None
    if gemini_check:
        (
            gemini_findings_by_idx, gemini_cost_usd, gemini_truncated, gemini_errored, gemini_error_detail,
        ) = await _run_gemini_chunks(
            item_chunks, checks, extra_instructions, lang, source_lang, semaphore,
        )
        gemini_findings_by_idx = _resolve_repeated_findings(gemini_findings_by_idx, relevant_rows)
        gemini_findings_by_idx, _ = _extract_register_values(gemini_findings_by_idx)
        # Same register-contamination fix as the calibration_debug pass
        # above (see that block's comment) — Gemini gets sent the exact
        # same prompt, including the register-instructions block when
        # "register" is selected, so it can return the same register_mixed
        # noise, which isn't what this second-opinion pass is for.
        gemini_findings_by_idx = {
            idx: [f for f in findings if f.get("type") != REGISTER_MIXED_TYPE]
            for idx, findings in gemini_findings_by_idx.items()
        }
        gemini_findings_by_idx = {idx: fs for idx, fs in gemini_findings_by_idx.items() if fs}

    out = []
    for idx, row in enumerate(relevant_rows):
        src = row["values"].get(source_lang, "")
        tgt = row["values"].get(lang, "")
        findings = run_rule_checks(src, tgt, checks, max_length=row["max_length"], lang_code=lang)
        findings += ai_findings_by_idx.get(idx, [])
        if calibration_debug:
            findings += _mark_calibration_debug_findings(ai_findings_by_idx_relaxed.get(idx, []))
        if gemini_check:
            findings += _mark_gemini_findings(gemini_findings_by_idx.get(idx, []))
        if findings:
            out.append({
                "excel_row": row["excel_row"],
                "context": row["context"],
                "source": src,
                "translation": tgt,
                "findings": findings,
            })
    if truncated:
        # The model's response for this language got cut off mid-array —
        # some rows may never have been checked by it at all. Surfaced as
        # its own synthetic entry rather than silently showing whatever
        # partial findings survived as if they were the complete picture
        # (see Александр's "incomplete report" on a large Spanish upload).
        out.append({
            "excel_row": 0,
            "context": "⚠ Системное предупреждение",
            "source": "",
            "translation": "",
            "findings": [_truncation_warning()],
        })
    if calibration_debug and debug_truncated:
        out.append({
            "excel_row": 0,
            "context": "⚠ Системное предупреждение",
            "source": "",
            "translation": "",
            "findings": [{
                "type": "system",
                "severity": "low",
                "message": (
                    "🔬 Тестовый прогон с мягкой калибровкой для этого языка был обрезан из-за большого "
                    "объёма — часть строк могла остаться непроверенной именно в тестовом (не обычном) "
                    "проходе."
                ),
            }],
        })
    if gemini_check and gemini_truncated:
        out.append({
            "excel_row": 0,
            "context": "⚠ Системное предупреждение",
            "source": "",
            "translation": "",
            "findings": [{
                "type": "system",
                "severity": "low",
                "message": (
                    "🌐 Проверка через Gemini для этого языка была обрезана из-за большого объёма — часть "
                    "строк могла остаться непроверенной именно этой дополнительной проверкой. На обычную "
                    "проверку через Claude это не повлияло."
                ),
            }],
        })
    if gemini_check and gemini_errored:
        # Added 2026-09-22 after a real Railway run only showed the generic
        # "ошибка сети, ключа или модели" text with no way to tell which of
        # the three it actually was — now names the specific reason
        # (_gemini_error_detail) right in the report itself, so Александр
        # doesn't need to check Railway's own logs to know what to fix.
        _gemini_reason = gemini_error_detail or "ошибка сети, ключа или модели"
        out.append({
            "excel_row": 0,
            "context": "⚠ Системное предупреждение",
            "source": "",
            "translation": "",
            "findings": [{
                "type": "system",
                "severity": "low",
                "message": (
                    f"🌐 Проверка через Gemini для этого языка не выполнилась: {_gemini_reason}. Не повлияло "
                    f"на обычную проверку через Claude, но эта дополнительная проверка не сработала."
                ),
            }],
        })
    if "register" in checks:
        by_excel_row = {relevant_rows[idx]["excel_row"]: v for idx, v in register_values_by_idx.items()}
        texts_by_excel_row = {
            relevant_rows[idx]["excel_row"]: relevant_rows[idx]["values"].get(lang, "")
            for idx in register_values_by_idx
        }
        block = _register_summary_block(build_register_report(by_excel_row, texts_by_excel_row))
        if block is not None:
            out.append(block)
    return out, cost_usd, debug_cost_usd, gemini_cost_usd


async def run_multi_check(
    sheets: list[dict],
    source_lang: str,
    checks: list[str],
    extra_instructions: str = "",
    target_langs_filter: set[str] | None = None,
    calibration_debug: bool = False,
    gemini_check: bool = False,
) -> dict:
    """
    Each target language gets its own AI call, so the prompt for e.g.
    "es-mx" never carries the other 34 languages' rows.

    target_langs_filter: when given, only these languages are checked even
    if the file has more columns — lets a manager check a subset of a
    large upload instead of every language every time.

    calibration_debug: see _check_language_for_sheet's own comment — runs
    every language's AI check a SECOND time with a loosened confidence bar
    (same model, same rows), and adds whatever that extra pass catches to
    the report as clearly marked "🔬" findings, so the two passes can be
    compared side by side. This roughly doubles the AI cost of the run
    (summary["calibration_debug_cost_usd"] shows exactly how much of the
    total came from this extra pass) — off by default, only for a
    deliberate one-off comparison, never for routine checking.

    gemini_check: see _check_language_for_sheet's own comment and
    app.gemini_client — runs every language's AI check ALSO through Google
    Gemini (same prompt/calibration, different model provider) and adds
    whatever it catches as "🌐"-tagged findings, counted as real findings
    (unlike calibration_debug's test-only ones — see _mark_gemini_findings)
    since a blind test showed it independently catches real errors ours
    misses. Also roughly doubles AI cost (summary["gemini_cost_usd"]) —
    off by default, opt-in per run.
    """
    semaphore = asyncio.Semaphore(AI_CONCURRENCY)
    result_sheets = []
    total_findings = 0
    total_debug_findings = 0
    total_gemini_findings = 0
    total_rows_checked = 0
    total_cost_usd = 0.0
    total_debug_cost_usd = 0.0
    total_gemini_cost_usd = 0.0

    for sheet in sheets:
        target_langs = [l for l in sheet["languages"] if l != source_lang]
        if target_langs_filter is not None:
            target_langs = [l for l in target_langs if _lang_selected(l, target_langs_filter)]
        tasks = [
            _check_language_for_sheet(
                sheet, lang, source_lang, checks, extra_instructions, semaphore,
                calibration_debug=calibration_debug, gemini_check=gemini_check,
            )
            for lang in target_langs
        ]
        per_lang_results = await asyncio.gather(*tasks) if tasks else []

        dup_cols = sheet.get("duplicate_language_columns") or {}
        languages_out = {}
        for lang, (findings_list, lang_cost, lang_debug_cost, lang_gemini_cost) in zip(target_langs, per_lang_results):
            findings_list = _apply_duplicate_language_warnings(
                findings_list, lang, dup_cols, source_lang, show_source_warning=(lang == target_langs[0]),
            )
            languages_out[lang] = findings_list
            total_findings += _count_real_findings(findings_list)
            total_debug_findings += sum(
                1 for row in findings_list for f in row["findings"] if f.get("calibration_debug")
            )
            total_gemini_findings += sum(
                1 for row in findings_list for f in row["findings"] if f.get("gemini_check")
            )
            total_cost_usd += lang_cost + lang_debug_cost + lang_gemini_cost
            total_debug_cost_usd += lang_debug_cost
            total_gemini_cost_usd += lang_gemini_cost

        total_rows_checked += len(sheet["rows"])
        result_sheets.append({
            "sheet_name": sheet["sheet_name"],
            "source_lang": source_lang,
            "languages_checked": target_langs,
            "languages": languages_out,
            "unrecognized_columns": sheet.get("unrecognized_columns", []),
        })

    summary = {
        "sheets": len(result_sheets),
        "rows_checked": total_rows_checked,
        "languages_checked": sorted({l for s in result_sheets for l in s["languages_checked"]}),
        "total_findings": total_findings,
        "cost_usd": total_cost_usd,
    }
    if calibration_debug:
        summary["calibration_debug_cost_usd"] = total_debug_cost_usd
        summary["calibration_debug_findings"] = total_debug_findings
    if gemini_check:
        summary["gemini_cost_usd"] = total_gemini_cost_usd
        summary["gemini_findings"] = total_gemini_findings
    return {"sheets": result_sheets, "summary": summary}


# ------------------------------------------------- large jobs: batch mode ---
# See BATCH_THRESHOLD_CHARS above. Instead of awaiting every language's AI
# call directly (run_multi_check), a large job is prepared as a "skeleton"
# (rule-based findings, computed instantly and for free) plus one Anthropic
# Message Batch request per CHUNK of each language's rows (see
# MAX_ROWS_PER_AI_CALL — same chunking as the live path, same reason); once
# that batch finishes — polled from app.main — finalize_batch_results merges
# the AI findings back in to produce the exact same {"sheets": [...],
# "summary": {...}} shape as run_multi_check, so the frontend doesn't need
# to know which path ran.

def estimate_check_volume(
    sheets: list[dict], source_lang: str, target_langs_filter: set[str] | None = None
) -> int:
    """Rough proxy (total characters, summed across every checkable row ×
    every target language) for how expensive/slow a synchronous run would
    be. Not exact — a ballpark is all that's needed to pick a processing
    mode. Respects target_langs_filter so checking only a handful of a
    file's languages doesn't get pushed into the slow queue on the
    strength of languages that won't even be checked this run."""
    total = 0
    for sheet in sheets:
        target_langs = [l for l in sheet["languages"] if l != source_lang]
        if target_langs_filter is not None:
            target_langs = [l for l in target_langs if _lang_selected(l, target_langs_filter)]
        for row in sheet["rows"]:
            src = row["values"].get(source_lang, "")
            for lang in target_langs:
                tgt = row["values"].get(lang, "")
                if not src.strip() and not tgt.strip():
                    continue
                if _is_do_not_translate(tgt):
                    continue
                total += len(src) + len(tgt)
    return total


def build_batch_plan(
    sheets: list[dict],
    source_lang: str,
    checks: list[str],
    extra_instructions: str = "",
    target_langs_filter: set[str] | None = None,
) -> tuple[list[dict], dict]:
    """Prepares everything needed to submit one Anthropic Message Batch
    covering every (sheet, target language, row chunk) triple in this
    upload, plus a JSON-serializable "skeleton" — already-computed
    rule-based findings — to merge the AI results into later via
    finalize_batch_results.

    Returns (batch_requests, skeleton). batch_requests is a list of
    {"custom_id": str, "prompt": str} ready for
    claude_client.create_message_batch; it can be empty if no AI check
    types were selected at all, in which case there's nothing to submit and
    finalize_batch_results(skeleton, {}) is already the final answer.
    """
    requests: list[dict] = []
    skeleton_sheets = []

    for s_idx, sheet in enumerate(sheets):
        target_langs = [l for l in sheet["languages"] if l != source_lang]
        if target_langs_filter is not None:
            target_langs = [l for l in target_langs if _lang_selected(l, target_langs_filter)]
        languages_skeleton = {}

        for lang_idx, lang in enumerate(target_langs):
            relevant_rows = []
            ai_items = []
            for row in sheet["rows"]:
                src = row["values"].get(source_lang, "")
                tgt = row["values"].get(lang, "")
                if not src.strip() and not tgt.strip():
                    continue
                if _is_do_not_translate(tgt):
                    continue
                relevant_rows.append(row)
                ai_items.append({"context": row["context"], "source": src, "translation": tgt})

            # Rule-based findings are free and instant — compute them now
            # rather than waiting on the batch for them too.
            base_rows = []
            for row in relevant_rows:
                src = row["values"].get(source_lang, "")
                tgt = row["values"].get(lang, "")
                findings = run_rule_checks(src, tgt, checks, max_length=row["max_length"], lang_code=lang)
                base_rows.append({
                    "excel_row": row["excel_row"],
                    "context": row["context"],
                    "source": src,
                    "translation": tgt,
                    "findings": findings,
                })

            # Built from the sheet/position/chunk index only, never from
            # `lang` itself — Anthropic's Batches API requires custom_id to
            # match ^[a-zA-Z0-9_-]{1,64}$, but a language code comes
            # straight from a column header in whatever file gets uploaded
            # and can't be trusted to satisfy that (e.g. a Cyrillic
            # character that looks identical to a Latin one, from a
            # copy-pasted "fr-CI" header, is enough to make Anthropic
            # reject the WHOLE batch — every language in it, not just the
            # bad one — with a 400). Rebuilt this way, custom_id is always
            # safe regardless of what's in the file.
            model = _model_for_lang(lang)
            # See MAX_ROWS_PER_AI_CALL — one Message Batch REQUEST per
            # chunk of this language's rows, not one for the whole
            # language, for the same per-row-attention reason as the live
            # path (_check_language_for_sheet) — including the smaller
            # MAX_ROWS_PER_AI_CALL_HARD chunk size for a hard language, so
            # a large upload doesn't get worse per-row attention than a
            # small one just because it happened to cross BATCH_THRESHOLD_CHARS.
            chunks_skeleton = []
            offset = 0
            for chunk_idx, chunk_items in enumerate(_chunk_list(ai_items, _chunk_size_for_lang(lang))):
                custom_id = f"s{s_idx}-t{lang_idx}-c{chunk_idx}"
                prompt, number_to_index = build_batch_prompt(
                    chunk_items, checks, extra_instructions, lang, source_lang,
                )
                if prompt is not None:
                    requests.append({"custom_id": custom_id, "prompt": prompt, "model": model})
                chunks_skeleton.append({
                    "custom_id": custom_id if prompt is not None else None,
                    "number_to_index": {str(k): v for k, v in number_to_index.items()},
                    # Where this chunk's local item indices (0-based, reset
                    # per chunk) land in the language's own `rows` list —
                    # finalize_batch_results adds this back before doing
                    # anything else with the model's response.
                    "row_offset": offset,
                })
                offset += len(chunk_items)

            languages_skeleton[lang] = {
                # Needed later by finalize_batch_results to price this
                # language's usage at the right per-token rate.
                "model": model,
                "chunks": chunks_skeleton,
                "rows": base_rows,
            }

        skeleton_sheets.append({
            "sheet_name": sheet["sheet_name"],
            "target_langs": target_langs,
            "languages": languages_skeleton,
            "unrecognized_columns": sheet.get("unrecognized_columns", []),
            "duplicate_language_columns": sheet.get("duplicate_language_columns", {}),
            "row_count": len(sheet["rows"]),
        })

    # checks rides along in the skeleton so finalize_batch_results (called
    # later, sometimes in a completely different request once the
    # Anthropic batch has ended) can still filter the model's response to
    # only what was actually asked for.
    skeleton = {"sheets": skeleton_sheets, "source_lang": source_lang, "checks": checks}
    return requests, skeleton


def finalize_batch_results(skeleton: dict, ai_results_by_custom_id: dict[str, dict]) -> dict:
    """Merges AI findings (once the Anthropic batch has ended) into the
    rule-based skeleton from build_batch_plan, producing the same
    {"sheets": [...], "summary": {...}} shape run_multi_check returns.

    Each language may have several chunks (see MAX_ROWS_PER_AI_CALL) —
    every chunk's custom_id is looked up and merged independently, with its
    own local item indices shifted back by its "row_offset" before doing
    anything else with them.

    ai_results_by_custom_id: {custom_id: {"text": str | None, "usage": dict}}
    — see claude_client.get_batch_results. Missing/empty entries (e.g. no
    batch was actually submitted) simply contribute no findings and no cost."""
    result_sheets = []
    total_findings = 0
    total_rows_checked = 0
    total_cost_usd = 0.0

    for sheet in skeleton["sheets"]:
        languages_out = {}
        dup_cols = sheet.get("duplicate_language_columns") or {}
        first_target_lang = (sheet.get("target_langs") or [None])[0]
        for lang, lang_skel in sheet["languages"].items():
            # One chunk's problem (missing/errored/truncated — see below)
            # never throws away another chunk's perfectly good findings;
            # each chunk is its own independent AI call (see
            # MAX_ROWS_PER_AI_CALL/build_batch_plan). At most one warning is
            # still shown per language, worst-first, so the manager isn't
            # confused by a wall of near-identical warnings on a language
            # that got split into many chunks.
            ai_grouped: dict[int, list[dict]] = {}
            missing_any = False
            errored_reason = None
            truncated_any = False
            chunks = lang_skel.get("chunks")
            if chunks is None:
                # Backward compatibility: a Message Batch submitted BEFORE
                # this chunking change (2026-09-17) was persisted to the
                # database at submission time (see multi_check in
                # app.main) with the OLD skeleton shape — one custom_id/
                # number_to_index pair directly on the language, no
                # "chunks" list at all. Any such upload still "processing"
                # across this deploy must still resolve correctly once
                # its batch ends — without this, `lang_skel.get("chunks",
                # [])` would silently return [], and that language would
                # come back looking completely clean, discarding every
                # real AI finding it already paid for. Wrapping the old
                # shape as a single one-chunk list re-uses the exact same
                # merge logic below for it.
                old_custom_id = lang_skel.get("custom_id")
                chunks = (
                    [{
                        "custom_id": old_custom_id,
                        "number_to_index": lang_skel.get("number_to_index", {}),
                        "row_offset": 0,
                    }]
                    if old_custom_id is not None else []
                )
            for chunk in chunks:
                custom_id = chunk["custom_id"]
                if custom_id is None:
                    continue
                offset = chunk["row_offset"]
                ai_result = ai_results_by_custom_id.get(custom_id)
                if ai_result is None:
                    # Expected result never showed up in the batch's .jsonl
                    # at all — same class of silent data loss as an
                    # errored/expired request, so it gets the same warning
                    # rather than quietly counting as "nothing found".
                    missing_any = True
                    continue
                if ai_result.get("result_type") != "succeeded":
                    errored_reason = errored_reason or (ai_result.get("result_type") or "неизвестная ошибка")
                else:
                    raw = parse_json_array(ai_result.get("text"))
                    # JSON round-trips dict keys as strings — restore int keys.
                    number_to_index = {int(k): v for k, v in chunk["number_to_index"].items()}
                    chunk_grouped = group_batch_findings(raw, number_to_index)
                    # Guarantees the model's response never smuggles in a check
                    # type the manager didn't ask for, even if it ignored the
                    # prompt's instruction to stick to the requested list.
                    chunk_grouped = {
                        idx: _filter_findings_by_checks(fs, skeleton.get("checks", []))
                        for idx, fs in chunk_grouped.items()
                    }
                    for idx, fs in chunk_grouped.items():
                        # "_also_idx" (see group_batch_findings) is still
                        # local to this chunk here — shift it the same way
                        # as the top-level index, or _resolve_repeated_
                        # findings below resolves it against the wrong rows.
                        remapped_fs = []
                        for f in fs:
                            if "_also_idx" in f:
                                f = {**f, "_also_idx": [offset + i for i in f["_also_idx"]]}
                            remapped_fs.append(f)
                        ai_grouped[offset + idx] = remapped_fs
                    if ai_result.get("stop_reason") == "max_tokens":
                        # Response got cut off mid-array — some rows in
                        # THIS chunk may never have been checked by the AI
                        # at all (see Александр's "incomplete Spanish
                        # report") — the chunk's own findings up to the cut
                        # point are still kept, just flagged.
                        truncated_any = True
                total_cost_usd += _usage_cost(lang_skel.get("model", ""), ai_result.get("usage"), batch=True)

            if missing_any:
                warning_finding = _ai_failure_warning("результат не получен")
            elif errored_reason:
                warning_finding = _ai_failure_warning(errored_reason)
            elif truncated_any:
                warning_finding = _truncation_warning()
            else:
                warning_finding = None

            ai_grouped = _resolve_repeated_findings(ai_grouped, lang_skel["rows"])
            ai_grouped, register_values_by_idx = _extract_register_values(ai_grouped)

            findings_list = []
            for idx, row in enumerate(lang_skel["rows"]):
                findings = list(row["findings"]) + ai_grouped.get(idx, [])
                if findings:
                    findings_list.append({
                        "excel_row": row["excel_row"],
                        "context": row["context"],
                        "source": row["source"],
                        "translation": row["translation"],
                        "findings": findings,
                    })
            if warning_finding is not None:
                findings_list.append({
                    "excel_row": 0,
                    "context": "⚠ Системное предупреждение",
                    "source": "",
                    "translation": "",
                    "findings": [warning_finding],
                })
            if "register" in skeleton.get("checks", []):
                by_excel_row = {
                    lang_skel["rows"][idx]["excel_row"]: v for idx, v in register_values_by_idx.items()
                }
                texts_by_excel_row = {
                    lang_skel["rows"][idx]["excel_row"]: lang_skel["rows"][idx]["translation"]
                    for idx in register_values_by_idx
                }
                block = _register_summary_block(build_register_report(by_excel_row, texts_by_excel_row))
                if block is not None:
                    findings_list.append(block)
            findings_list = _apply_duplicate_language_warnings(
                findings_list, lang, dup_cols, skeleton["source_lang"],
                show_source_warning=(lang == first_target_lang),
            )
            languages_out[lang] = findings_list
            total_findings += _count_real_findings(findings_list)

        total_rows_checked += sheet["row_count"]
        result_sheets.append({
            "sheet_name": sheet["sheet_name"],
            "source_lang": skeleton["source_lang"],
            "languages_checked": sheet["target_langs"],
            "languages": languages_out,
            "unrecognized_columns": sheet["unrecognized_columns"],
        })

    summary = {
        "sheets": len(result_sheets),
        "rows_checked": total_rows_checked,
        "languages_checked": sorted({l for s in result_sheets for l in s["languages_checked"]}),
        "total_findings": total_findings,
        "cost_usd": total_cost_usd,
    }
    return {"sheets": result_sheets, "summary": summary}


async def submit_multi_check_batch(requests: list[dict]) -> str | None:
    return await create_message_batch(requests)


async def cancel_multi_check_batch(batch_id: str) -> None:
    """Best-effort cancel for a manager who no longer wants to wait for (or
    pay for) a still-processing upload — e.g. they want to switch to
    "Срочно" instead, or simply changed their mind. Anthropic may reject
    this (most likely because the batch had already ended right as the
    manager clicked cancel) — that's fine, the caller is deleting its own
    record either way, so a failure here should never block that."""
    try:
        await cancel_message_batch(batch_id)
    except Exception:
        pass


def _batch_progress(status: dict) -> dict:
    """Turns Anthropic's own request_counts ({"processing": n, "succeeded":
    n, "errored": n, "canceled": n, "expired": n}) into a simple {"done",
    "total"} the UI can show as a real progress readout. Anthropic doesn't
    publish an ETA for a batch job, so a made-up time estimate would just be
    a guess — this is the one number we actually know is true."""
    counts = status.get("request_counts") or {}
    total = sum(counts.values())
    done = total - counts.get("processing", 0)
    return {"done": done, "total": total}


async def try_finalize_batch(batch_id: str, skeleton: dict) -> tuple[dict | None, dict]:
    """Returns (finalized_results, progress). finalized_results is the
    completed results dict once the Anthropic batch has ended, otherwise
    None (still processing — caller should try again later). progress is
    always {"done": int, "total": int} from Anthropic's request_counts, so
    the caller can surface real progress even while still waiting."""
    status = await get_batch_status(batch_id)
    progress = _batch_progress(status)
    if status.get("processing_status") != "ended":
        return None, progress
    results_url = status.get("results_url")
    ai_results_by_custom_id = await get_batch_results(results_url) if results_url else {}
    return finalize_batch_results(skeleton, ai_results_by_custom_id), progress


def _plural_ru(n: int, one: str, few: str, many: str) -> str:
    """Standard Russian count-noun pluralization — mirrors the frontend's
    pluralRu (lang.ts), kept as a separate copy since this side is Python."""
    mod10, mod100 = n % 10, n % 100
    if 11 <= mod100 <= 14:
        return many
    if mod10 == 1:
        return one
    if 2 <= mod10 <= 4:
        return few
    return many


def _format_minutes_ru(minutes: float) -> str:
    """Mirrors the frontend's formatElapsedMinutesRu (lang.ts) — used in the
    downloadable report's summary line, below."""
    whole = round(minutes)
    if whole < 1:
        return "меньше минуты"
    return f"{whole} {_plural_ru(whole, 'минута', 'минуты', 'минут')}"


def build_report_workbook(
    filename: str, source_lang: str, results: dict, duration_minutes: float | None = None
) -> bytes:
    """Builds a downloadable .xlsx with one row per finding. duration_minutes
    (how long the check itself took, end to end) is optional — omitted
    entirely for an older record that predates this being tracked, rather
    than showing a misleading "0 минут"."""
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "QA Findings"
    if duration_minutes is not None:
        ws.append([f"Проверка «{filename}» ({source_lang}) — заняла {_format_minutes_ru(duration_minutes)}"])
        ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=9)
        ws.cell(row=1, column=1).font = openpyxl.styles.Font(bold=True)
        ws.append([])  # spacer row before the header
    ws.append(["Лист", "Строка в файле", "Контекст", "Язык", "Серьёзность", "Тип", "Проблема", "Источник", "Перевод"])
    # Captured AFTER the append above (not computed from the pre-append
    # max_row) — an appended blank spacer row still advances openpyxl's
    # internal row cursor even though it holds no cells, so computing this
    # beforehand pointed one row too early and left the header itself out
    # of the filter/freeze range below.
    header_row = ws.max_row
    for col_idx, width in enumerate([18, 14, 28, 8, 12, 14, 50, 40, 40], start=1):
        ws.column_dimensions[openpyxl.utils.get_column_letter(col_idx)].width = width

    for sheet in results.get("sheets", []):
        for lang, findings_list in sheet.get("languages", {}).items():
            for item in findings_list:
                for f in item["findings"]:
                    ws.append([
                        sheet["sheet_name"],
                        item["excel_row"],
                        item["context"],
                        lang,
                        f.get("severity", ""),
                        f.get("type", ""),
                        f.get("message", ""),
                        item["source"],
                        item["translation"],
                    ])

    # Turns on Excel's own column filter dropdowns on the header row — lets
    # Александр filter to just one "Язык" (or any other column) using the
    # filter control Excel already gives him, defaulting to showing
    # everything exactly as it does today. Only makes sense once there's at
    # least one data row below the header.
    if ws.max_row > header_row:
        ws.auto_filter.ref = f"A{header_row}:I{ws.max_row}"
        ws.freeze_panes = f"A{header_row + 1}"

    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()
