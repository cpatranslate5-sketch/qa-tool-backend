"""
Word / PowerPoint / JSON → the same table the Excel check reads
(2026-10-05, Александр: «можно ли проверять документы других форматов?»).

Every format is turned into an .xlsx workbook — one header row with
language codes (plus a «Context»/«Key» column), one row per text unit —
and from there the usual check runs unchanged: language detection,
confirmation, the AI pipeline, reports, translator links.

Supported:
  * one .docx with a table (or several) — original and translation in
    columns; headers may be language codes, «Оригинал»/«Перевод», or none;
  * two .docx files — original and translation; paragraphs (and table
    cells) are paired automatically, by order and length;
  * two .pptx files — paired slide by slide, text box by text box;
  * two .json files — paired by key; or one .json whose top-level keys are
    languages ({"ru": {...}, "kk": {...}}).
"""
from __future__ import annotations

import io
import json
import math
import re

import openpyxl

from app.excel_multi import parse_workbook

FORMATS = (".xlsx", ".docx", ".pptx", ".json")


class ConvertError(ValueError):
    """A problem the manager can fix — the message is shown as is."""


def _ext(name: str) -> str:
    name = (name or "").lower()
    for e in FORMATS:
        if name.endswith(e):
            return e
    return ""


def _clean(text: str | None) -> str:
    return re.sub(r"[ \t ]+", " ", (text or "")).strip()


def _workbook(sheets: list[tuple[str, list[str], list[list[str]]]]) -> bytes:
    wb = openpyxl.Workbook()
    wb.remove(wb.active)
    used: set[str] = set()
    for title, header, rows in sheets:
        t = re.sub(r"[\[\]:*?/\\]", " ", title)[:31] or "Лист"
        base, k = t, 2
        while t.lower() in used:
            t = f"{base[:27]} ({k})"
            k += 1
        used.add(t.lower())
        ws = wb.create_sheet(t)
        ws.append(header)
        for r in rows:
            ws.append(r)
    out = io.BytesIO()
    wb.save(out)
    return out.getvalue()


# ------------------------------------------------------------- alignment ---
_NUM_RE = re.compile(r"\d+")


def _pair_cost(a: str, b: str, ratio: float) -> float:
    la, lb = len(a), len(b)
    cost = abs(math.log((lb + 8) / (la * ratio + 8))) * 4.0
    na, nb = set(_NUM_RE.findall(a)), set(_NUM_RE.findall(b))
    if na or nb:
        if na == nb:
            cost -= 1.0
        elif not (na & nb):
            cost += 1.5
    return cost


def align(src: list[str], tgt: list[str]) -> list[tuple[list[int], list[int]]]:
    """Pairs two lists of text units (paragraphs) that are translations of
    each other: 1–1, 1–2, 2–1 and unmatched (1–0 / 0–1), by order, length
    and the numbers they contain. Returns [(src indexes, tgt indexes)]."""
    n, m = len(src), len(tgt)
    if n == 0 or m == 0:
        return [([i], []) for i in range(n)] + [([], [j]) for j in range(m)]
    total_s = sum(len(s) for s in src) or 1
    ratio = (sum(len(t) for t in tgt) or 1) / total_s
    if n == m and all(_pair_cost(s, t, ratio) < 3.0 for s, t in zip(src, tgt)):
        return [([i], [i]) for i in range(n)]
    band = max(40, int(0.15 * max(n, m)))
    INF = float("inf")
    cost: dict[tuple[int, int], float] = {(0, 0): 0.0}
    back: dict[tuple[int, int], tuple[int, int]] = {}
    SKIP, MERGE = 4.0, 1.5
    for i in range(n + 1):
        centre = i * m / n
        lo, hi = max(0, int(centre - band)), min(m, int(centre + band) + 1)
        for j in range(lo, hi + 1):
            if (i, j) == (0, 0):
                continue
            best, arg = INF, None
            for di, dj in ((1, 1), (1, 0), (0, 1), (2, 1), (1, 2)):
                pi, pj = i - di, j - dj
                if pi < 0 or pj < 0:
                    continue
                prev = cost.get((pi, pj))
                if prev is None:
                    continue
                if (di, dj) == (1, 1):
                    c = _pair_cost(src[pi], tgt[pj], ratio)
                elif (di, dj) == (2, 1):
                    c = _pair_cost(src[pi] + " " + src[pi + 1], tgt[pj], ratio) + MERGE
                elif (di, dj) == (1, 2):
                    c = _pair_cost(src[pi], tgt[pj] + " " + tgt[pj + 1], ratio) + MERGE
                else:
                    c = SKIP
                if prev + c < best:
                    best, arg = prev + c, (pi, pj)
            if arg is not None:
                cost[(i, j)] = best
                back[(i, j)] = arg
    if (n, m) not in cost:  # band too narrow for wildly different lengths
        return [([i], [i]) if i < m else ([i], []) for i in range(n)] + [([], [j]) for j in range(n, m)]
    pairs = []
    i, j = n, m
    while (i, j) != (0, 0):
        pi, pj = back[(i, j)]
        pairs.append((list(range(pi, i)), list(range(pj, j))))
        i, j = pi, pj
    pairs.reverse()
    return pairs


def _aligned_rows(src: list[tuple[str, str]], tgt: list[tuple[str, str]], warnings: list[str],
                  where: str = "") -> list[list[str]]:
    """src/tgt: [(context, text)] → rows [context, original, translation]."""
    pairs = align([t for _, t in src], [t for _, t in tgt])
    rows, lost_src, lost_tgt = [], [], []
    for si, ti in pairs:
        s_text = "\n".join(src[i][1] for i in si)
        t_text = "\n".join(tgt[j][1] for j in ti)
        ctx = src[si[0]][0] if si else (tgt[ti[0]][0] if ti else "")
        if si and not ti:
            lost_src.append(s_text)
        if ti and not si:
            lost_tgt.append(t_text)
            continue  # a translation with no original can't be checked
        rows.append([ctx, s_text, t_text])
    prefix = f"{where}: " if where else ""
    if lost_src:
        warnings.append(f"{prefix}для {len(lost_src)} фрагм. оригинала не нашлось пары в переводе "
                        f"(например: «{lost_src[0][:80]}») — они попадут в проверку с пустым переводом.")
    if lost_tgt:
        warnings.append(f"{prefix}{len(lost_tgt)} фрагм. перевода без пары в оригинале пропущены "
                        f"(например: «{lost_tgt[0][:80]}»).")
    return rows


# ------------------------------------------------------------------ Word ---
def _docx(data: bytes):
    try:
        import docx  # python-docx
        return docx.Document(io.BytesIO(data))
    except ImportError as e:
        raise ConvertError("На сервере не установлена поддержка Word (python-docx).") from e
    except Exception as e:
        raise ConvertError("Не удалось открыть документ Word — проверьте, что это файл .docx.") from e


def _table_rows(table) -> list[list[str]]:
    rows = []
    for row in table.rows:
        cells, seen = [], set()
        for cell in row.cells:
            if id(cell._tc) in seen:  # merged cell repeats itself
                continue
            seen.add(id(cell._tc))
            cells.append(_clean(cell.text))
        if any(cells):
            rows.append(cells)
    width = max((len(r) for r in rows), default=0)
    return [r + [""] * (width - len(r)) for r in rows]


def _docx_units(doc) -> list[tuple[str, str]]:
    """Paragraphs and table cells in document order: [(context, text)]."""
    from docx.table import Table
    from docx.text.paragraph import Paragraph

    units, heading, t_no = [], "", 0
    body = doc.element.body
    for child in body.iterchildren():
        tag = child.tag.rsplit("}", 1)[-1]
        if tag == "p":
            p = Paragraph(child, doc)
            text = _clean(p.text)
            if not text:
                continue
            style = (p.style.name if p.style is not None else "") or ""
            if style.lower().startswith(("heading", "заголовок", "title")):
                heading = text[:60]
            units.append((f"Раздел: {heading}" if heading else "", text))
        elif tag == "tbl":
            t_no += 1
            for r_no, row in enumerate(_table_rows(Table(child, doc)), start=1):
                for cell in row:
                    if cell:
                        units.append((f"Таблица {t_no}, строка {r_no}", cell))
    return units


_SRC_WORDS = ("оригинал", "original", "source", "исходн", "источник", "эталон")
_TGT_WORDS = ("перевод", "translation", "target", "переклад", "аударма", "tarjima", "тарҷума")


def _docx_tables_to_sheets(doc, source_lang: str, target_lang: str, alias_map, warnings) -> list:
    sheets = []
    for t_no, table in enumerate(doc.tables, start=1):
        rows = _table_rows(table)
        if len(rows) < 1 or len(rows[0]) < 2:
            continue
        title = f"Таблица {t_no}"
        # 1) the table's own header already names the languages?
        probe = _workbook([(title, rows[0], rows[1:])])
        try:
            parsed = parse_workbook(probe, alias_map)
        except Exception:
            parsed = []
        langs = {l for l in (parsed[0]["languages"] if parsed else []) if _LANGISH.match(l)}
        if len(langs) >= 2 and any(l.split("-")[0] == source_lang.split("-")[0] for l in langs):
            sheets.append((title, rows[0], rows[1:]))
            continue
        # 2) «Оригинал» / «Перевод» headers, or none at all.
        if not target_lang:
            raise ConvertError(f"В таблице {t_no} не подписаны языки колонок — выберите язык перевода.")
        head = [h.lower() for h in rows[0]]
        src_col = next((i for i, h in enumerate(head) if any(w in h for w in _SRC_WORDS)), None)
        tgt_col = next((i for i, h in enumerate(head) if any(w in h for w in _TGT_WORDS)), None)
        data = rows[1:] if (src_col is not None or tgt_col is not None) else rows
        width = len(rows[0])
        textual = [i for i in range(width)
                   if sum(1 for r in data if r[i] and not re.fullmatch(r"[\d.,\s№#]+", r[i])) >= max(1, len(data) // 2)]
        if src_col is None or tgt_col is None:
            if len(textual) < 2:
                warnings.append(f"Таблица {t_no}: не нашлось двух колонок с текстом — пропущена.")
                continue
            src_col = src_col if src_col is not None else textual[0]
            tgt_col = tgt_col if tgt_col is not None else next(i for i in reversed(textual) if i != src_col)
            warnings.append(f"Таблица {t_no}: колонки определены автоматически — оригинал: колонка {src_col + 1}, "
                            f"перевод: колонка {tgt_col + 1}. Проверьте в таблице сопоставления.")
        others = [i for i in range(width) if i not in (src_col, tgt_col)]
        out = []
        for r in data:
            ctx = " | ".join(f"{rows[0][i]}: {r[i]}" if data is not rows and rows[0][i] else r[i]
                             for i in others if r[i])
            if r[src_col] or r[tgt_col]:
                out.append([ctx, r[src_col], r[tgt_col]])
        sheets.append((title, ["Context", source_lang, target_lang], out))
    return sheets


# ------------------------------------------------------------ PowerPoint ---
def _pptx_slides(data: bytes) -> list[list[tuple[str, str]]]:
    try:
        from pptx import Presentation
        from pptx.enum.shapes import MSO_SHAPE_TYPE
    except ImportError as e:
        raise ConvertError("На сервере не установлена поддержка PowerPoint (python-pptx).") from e
    try:
        prs = Presentation(io.BytesIO(data))
    except Exception as e:
        raise ConvertError("Не удалось открыть презентацию — проверьте, что это файл .pptx.") from e

    def shape_units(shape, ctx, out):
        if shape.shape_type == MSO_SHAPE_TYPE.GROUP:
            for s in shape.shapes:
                shape_units(s, ctx, out)
            return
        if getattr(shape, "has_table", False) and shape.has_table:
            for row in shape.table.rows:
                for cell in row.cells:
                    t = _clean(cell.text)
                    if t:
                        out.append((ctx, t))
            return
        if getattr(shape, "has_text_frame", False) and shape.has_text_frame:
            text = "\n".join(_clean(p.text) for p in shape.text_frame.paragraphs if _clean(p.text))
            if text:
                out.append((ctx, text))

    slides = []
    for n, slide in enumerate(prs.slides, start=1):
        units: list[tuple[str, str]] = []
        for shape in slide.shapes:
            shape_units(shape, f"Слайд {n}", units)
        if slide.has_notes_slide:
            notes = _clean(slide.notes_slide.notes_text_frame.text if slide.notes_slide.notes_text_frame else "")
            if notes:
                units.append((f"Слайд {n}, заметки", notes))
        slides.append(units)
    return slides


# ------------------------------------------------------------------ JSON ---
def _json(data: bytes, name: str):
    try:
        return json.loads(data.decode("utf-8-sig"))
    except Exception as e:
        raise ConvertError(f"Файл {name} — не корректный JSON.") from e


def _flatten(obj, prefix: str = "", out: dict | None = None) -> dict:
    out = {} if out is None else out
    if isinstance(obj, dict):
        for k, v in obj.items():
            _flatten(v, f"{prefix}.{k}" if prefix else str(k), out)
    elif isinstance(obj, list):
        for i, v in enumerate(obj):
            _flatten(v, f"{prefix}[{i}]", out)
    elif isinstance(obj, str):
        out[prefix] = obj
    return out


_LANGISH = re.compile(r"^[a-z]{2,3}([-_][a-z0-9]{2,4})?$", re.I)


# ------------------------------------------------------------------- main ---
def convert(files: list[tuple[str, bytes]], source_lang: str, target_lang: str,
            alias_map: dict | None = None) -> tuple[bytes, dict]:
    """→ (xlsx bytes, info {rows, sheets, warnings, mode})."""
    source_lang = (source_lang or "").strip().lower()
    target_lang = (target_lang or "").strip().lower()
    if not files:
        raise ConvertError("Файл не выбран.")
    if not source_lang:
        raise ConvertError("Сначала выберите язык оригинала.")
    exts = {_ext(n) for n, _ in files}
    if "" in exts:
        raise ConvertError("Поддерживаются файлы Excel (.xlsx), Word (.docx), PowerPoint (.pptx) и JSON (.json).")
    if len(files) > 2:
        raise ConvertError("Можно загрузить один файл или два: оригинал и перевод.")
    if len(files) == 2 and len(exts) != 1:
        raise ConvertError("Оригинал и перевод должны быть в одном формате.")
    ext = exts.pop()
    warnings: list[str] = []
    header = ["Context", source_lang, target_lang]

    if len(files) == 1:
        name, data = files[0]
        if ext == ".docx":
            sheets = _docx_tables_to_sheets(_docx(data), source_lang, target_lang, alias_map, warnings)
            if not sheets:
                raise ConvertError("В документе нет таблицы с оригиналом и переводом. Если перевод в отдельном "
                                   "файле — загрузите оригинал и перевод двумя файлами.")
            mode = "Word: таблица в документе"
        elif ext == ".json":
            obj = _json(data, name)
            if not (isinstance(obj, dict) and len(obj) >= 2 and sum(bool(_LANGISH.match(k)) for k in obj) >= 2):
                raise ConvertError("В JSON один язык. Загрузите два файла — оригинал и перевод — или один файл, "
                                   "где верхние ключи — языки ({\"ru\": {…}, \"kk\": {…}}).")
            langs = [k for k in obj if _LANGISH.match(k)]
            flat = {k: _flatten(obj[k]) for k in langs}
            src_key = next((k for k in langs if k.lower().split("-")[0].split("_")[0] == source_lang.split("-")[0]), None)
            if src_key is None:
                raise ConvertError(f"В JSON нет языка оригинала «{source_lang}».")
            order = [src_key] + [k for k in langs if k != src_key]
            rows = [[key] + [flat[k].get(key, "") for k in order] for key in flat[src_key]]
            sheets = [("JSON", ["Key"] + [k.replace("_", "-") for k in order], rows)]
            mode = "JSON: языки в одном файле"
        elif ext == ".pptx":
            raise ConvertError("Для презентации загрузите два файла: оригинал и перевод.")
        else:
            raise ConvertError("Excel-файл проверяется как обычно, без подготовки.")
    else:
        if not target_lang:
            raise ConvertError("Выберите язык перевода.")
        (s_name, s_data), (t_name, t_data) = files
        if ext == ".docx":
            src_u, tgt_u = _docx_units(_docx(s_data)), _docx_units(_docx(t_data))
            if not src_u:
                raise ConvertError("В оригинале не найден текст.")
            rows = _aligned_rows(src_u, tgt_u, warnings)
            sheets = [("Документ", header, rows)]
            mode = "Word: оригинал и перевод"
        elif ext == ".pptx":
            s_sl, t_sl = _pptx_slides(s_data), _pptx_slides(t_data)
            if len(s_sl) != len(t_sl):
                warnings.append(f"В оригинале {len(s_sl)} слайд(ов), в переводе — {len(t_sl)}; "
                                "слайды сопоставлены по порядку.")
            rows = []
            for k in range(len(s_sl)):
                rows += _aligned_rows(s_sl[k], t_sl[k] if k < len(t_sl) else [], warnings, f"Слайд {k + 1}")
            sheets = [("Презентация", header, rows)]
            mode = "PowerPoint: оригинал и перевод"
        elif ext == ".json":
            s_flat, t_flat = _flatten(_json(s_data, s_name)), _flatten(_json(t_data, t_name))
            rows = [[k, v, t_flat.get(k, "")] for k, v in s_flat.items()]
            missing = [k for k in s_flat if k not in t_flat]
            extra = [k for k in t_flat if k not in s_flat]
            if missing:
                warnings.append(f"В переводе нет {len(missing)} ключ(ей) оригинала (например: {missing[0]}).")
            if extra:
                warnings.append(f"В переводе {len(extra)} лишн. ключ(ей), которых нет в оригинале (например: {extra[0]}).")
            sheets = [("JSON", ["Key", source_lang, target_lang], rows)]
            mode = "JSON: оригинал и перевод"
        else:
            raise ConvertError("Два Excel-файла не нужны — оригинал и перевод должны быть колонками одного файла.")

    total = sum(len(r) for _, _, r in sheets)
    if total == 0:
        raise ConvertError("Не нашлось ни одной строки текста для проверки.")
    return _workbook(sheets), {"mode": mode, "rows": total, "sheets": len(sheets), "warnings": warnings}
